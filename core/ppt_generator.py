from __future__ import annotations

import math
import os
import re
import tempfile
import time
from functools import lru_cache
from pathlib import Path
from typing import Callable

from PIL import (
    Image,
    ImageOps,
    UnidentifiedImageError,
)
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

from core.heif_support import (
    is_heif_path,
    open_pillow_image,
    register_heif_support,
)
from core.group_layout import resolve_title_vertical_geometry


register_heif_support()


PAGE_SIZES = {
    "16:9": (33.867, 19.05),
    "4:3": (25.4, 19.05),
    "A4横版": (29.7, 21.0),
    "A4竖版": (21.0, 29.7),
}


PPT_EXPORT_DPI = 480
PIXELS_PER_CM = PPT_EXPORT_DPI / 2.54

# python-pptx可以直接写入的图像格式。
# MPO、WEBP、HEIF等格式需要先转换为高分辨率PNG。
PPT_DIRECT_IMAGE_FORMATS = frozenset(
    {
        "BMP",
        "GIF",
        "JPEG",
        "PNG",
        "TIFF",
        "WMF",
    }
)


@lru_cache(maxsize=4096)
def _source_image_format(
    src_path: str,
) -> str:
    """
    检测图片的真实编码格式，而不是只看扩展名。

    手机照片可能以 .jpg 结尾，但内部实际格式是 MPO。
    python-pptx会拒绝直接写入这种图片。
    """
    try:
        with open_pillow_image(
            src_path
        ) as image:
            return str(
                image.format or ""
            ).upper()
    except Exception:
        return ""


def _source_requires_ppt_conversion(
    src_path: str,
) -> bool:
    image_format = _source_image_format(
        str(src_path)
    )

    return bool(
        not image_format
        or image_format
        not in PPT_DIRECT_IMAGE_FORMATS
    )


def _normalize_crop(crop: dict | None) -> dict:
    crop = dict(crop or {})

    x = max(
        0.0,
        min(1.0, float(crop.get("x", 0.0))),
    )
    y = max(
        0.0,
        min(1.0, float(crop.get("y", 0.0))),
    )
    width = max(
        0.01,
        min(1.0 - x, float(crop.get("width", 1.0))),
    )
    height = max(
        0.01,
        min(1.0 - y, float(crop.get("height", 1.0))),
    )

    return {
        "enabled": bool(crop.get("enabled", False)),
        "x": x,
        "y": y,
        "width": width,
        "height": height,
    }


def _crop_is_modified(crop: dict | None) -> bool:
    crop = _normalize_crop(crop)

    return bool(
        crop["enabled"]
        or crop["x"] > 0.0001
        or crop["y"] > 0.0001
        or crop["width"] < 0.9999
        or crop["height"] < 0.9999
    )


def _transform_source(
    source: Image.Image,
    transform: dict,
) -> Image.Image:
    crop = _normalize_crop(
        transform.get("crop")
    )

    if _crop_is_modified(crop):
        left = round(source.width * crop["x"])
        top = round(source.height * crop["y"])
        right = round(
            source.width
            * (crop["x"] + crop["width"])
        )
        bottom = round(
            source.height
            * (crop["y"] + crop["height"])
        )

        source = source.crop(
            (
                max(0, left),
                max(0, top),
                min(source.width, max(left + 1, right)),
                min(source.height, max(top + 1, bottom)),
            )
        )

    rotation = int(
        transform.get("rotation", 0)
    ) % 360
    flip_h = bool(
        transform.get("flip_h", False)
    )
    flip_v = bool(
        transform.get("flip_v", False)
    )

    if flip_h:
        source = ImageOps.mirror(source)

    if flip_v:
        source = ImageOps.flip(source)

    if rotation:
        source = source.rotate(
            -rotation,
            expand=True,
            resample=Image.Resampling.BICUBIC,
        )

    return source


def _source_has_exif_rotation(
    src_path: str,
) -> bool:
    try:
        with open_pillow_image(src_path) as image:
            orientation = image.getexif().get(
                274,
                1,
            )
            return orientation not in {
                None,
                1,
            }
    except Exception:
        return False


def _requires_render(
    src_path: str,
    transform: dict,
) -> bool:
    """Return whether PPT needs a full-frame compatibility conversion."""
    return bool(
        is_heif_path(src_path)
        or _source_requires_ppt_conversion(
            src_path
        )
        or _source_has_exif_rotation(src_path)
        or int(transform.get("rotation", 0)) % 360
        or bool(transform.get("flip_h", False))
        or bool(transform.get("flip_v", False))
    )


def _image_pixel_size(
    src_path: str,
) -> tuple[int, int]:
    with open_pillow_image(src_path) as opened:
        try:
            opened.seek(0)
        except (EOFError, AttributeError):
            pass

        image = ImageOps.exif_transpose(
            opened
        )
        return (
            int(image.width),
            int(image.height),
        )


def _oriented_crop(
    crop: dict | None,
    *,
    rotation: int = 0,
    flip_h: bool = False,
    flip_v: bool = False,
) -> dict:
    """Map a pre-transform crop rectangle onto the converted full image."""
    result = _normalize_crop(crop)
    x = float(result["x"])
    y = float(result["y"])
    width = float(result["width"])
    height = float(result["height"])

    if flip_h:
        x = 1.0 - x - width
    if flip_v:
        y = 1.0 - y - height

    rotation = int(rotation) % 360
    if rotation == 90:
        x, y, width, height = (
            1.0 - y - height,
            x,
            height,
            width,
        )
    elif rotation == 180:
        x = 1.0 - x - width
        y = 1.0 - y - height
    elif rotation == 270:
        x, y, width, height = (
            y,
            1.0 - x - width,
            height,
            width,
        )

    return _normalize_crop(
        {
            "enabled": result["enabled"],
            "x": x,
            "y": y,
            "width": width,
            "height": height,
        }
    )


def _render_full_editable_source(
    src_path: str,
    transform: dict,
) -> str:
    """Convert a complete source image without discarding cropped pixels."""
    with open_pillow_image(src_path) as opened:
        try:
            opened.seek(0)
        except (EOFError, AttributeError):
            pass

        source = ImageOps.exif_transpose(opened).convert("RGBA")
        if bool(transform.get("flip_h", False)):
            source = ImageOps.mirror(source)
        if bool(transform.get("flip_v", False)):
            source = ImageOps.flip(source)

        rotation = int(transform.get("rotation", 0)) % 360
        if rotation:
            source = source.rotate(
                -rotation,
                expand=True,
                resample=Image.Resampling.BICUBIC,
            )

        temporary = tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".png",
        )
        temporary.close()
        source.save(temporary.name, "PNG", compress_level=1)
        return temporary.name


def _render_image_box(
    src_path: str,
    width_cm: float,
    height_cm: float,
    image_mode: str,
    transform: dict,
    *,
    export_dpi: int = PPT_EXPORT_DPI,
) -> str:
    """
    对有编辑操作的图片生成高分辨率临时PNG。

    关键修复：
    旧版先把原图压缩到目标框，再执行放大，zoom>1时会把
    已压缩图再次放大。新版先在原始像素上完成裁切、旋转、
    镜像和可见区域计算，最后只进行一次LANCZOS缩放。
    """
    pixels_per_cm = max(
        72,
        int(export_dpi),
    ) / 2.54

    target_w = max(
        320,
        round(width_cm * pixels_per_cm),
    )
    target_h = max(
        320,
        round(height_cm * pixels_per_cm),
    )

    zoom = max(
        0.10,
        min(
            10.00,
            float(transform.get("zoom", 1.0)),
        ),
    )
    offset_x = max(
        -50.0,
        min(
            50.0,
            float(transform.get("offset_x", 0.0)),
        ),
    )
    offset_y = max(
        -50.0,
        min(
            50.0,
            float(transform.get("offset_y", 0.0)),
        ),
    )

    with open_pillow_image(src_path) as opened:
        # MPO、GIF、TIFF等多帧图片统一使用主图/第1帧。
        try:
            opened.seek(0)
        except (EOFError, AttributeError):
            pass

        source = ImageOps.exif_transpose(
            opened
        ).convert("RGBA")
        source = _transform_source(
            source,
            transform,
        )

        if image_mode == "裁剪填充":
            base_scale = max(
                target_w / source.width,
                target_h / source.height,
            )
        else:
            base_scale = min(
                target_w / source.width,
                target_h / source.height,
            )

        scale = max(
            0.000001,
            base_scale * zoom,
        )
        scaled_w = source.width * scale
        scaled_h = source.height * scale

        pan_x = offset_x * target_w * 0.5
        pan_y = offset_y * target_h * 0.5
        image_left = (
            (target_w - scaled_w) / 2
            + pan_x
        )
        image_top = (
            (target_h - scaled_h) / 2
            + pan_y
        )

        dest_left = max(
            0,
            math.floor(image_left),
        )
        dest_top = max(
            0,
            math.floor(image_top),
        )
        dest_right = min(
            target_w,
            math.ceil(
                image_left + scaled_w
            ),
        )
        dest_bottom = min(
            target_h,
            math.ceil(
                image_top + scaled_h
            ),
        )

        canvas = Image.new(
            "RGBA",
            (target_w, target_h),
            (255, 255, 255, 0),
        )

        if (
            dest_right > dest_left
            and dest_bottom > dest_top
        ):
            source_left = max(
                0.0,
                (dest_left - image_left) / scale,
            )
            source_top = max(
                0.0,
                (dest_top - image_top) / scale,
            )
            source_right = min(
                float(source.width),
                (
                    dest_right - image_left
                ) / scale,
            )
            source_bottom = min(
                float(source.height),
                (
                    dest_bottom - image_top
                ) / scale,
            )

            crop_left = max(
                0,
                math.floor(source_left),
            )
            crop_top = max(
                0,
                math.floor(source_top),
            )
            crop_right = min(
                source.width,
                max(
                    crop_left + 1,
                    math.ceil(source_right),
                ),
            )
            crop_bottom = min(
                source.height,
                max(
                    crop_top + 1,
                    math.ceil(source_bottom),
                ),
            )

            visible_source = source.crop(
                (
                    crop_left,
                    crop_top,
                    crop_right,
                    crop_bottom,
                )
            )
            visible_target = visible_source.resize(
                (
                    dest_right - dest_left,
                    dest_bottom - dest_top,
                ),
                Image.Resampling.LANCZOS,
            )
            canvas.alpha_composite(
                visible_target,
                (
                    dest_left,
                    dest_top,
                ),
            )

    temporary = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".png",
    )
    temporary.close()

    canvas.save(
        temporary.name,
        "PNG",
        compress_level=1,
    )
    return temporary.name


def _add_editable_picture(
    slide,
    src_path: str,
    *,
    left_cm: float,
    top_cm: float,
    width_cm: float,
    height_cm: float,
    image_mode: str,
    transform: dict,
):
    """Embed a complete image and express the visible window as PPT crop."""
    source_w, source_h = _image_pixel_size(src_path)
    crop = _normalize_crop(transform.get("crop"))
    crop_x = float(crop["x"])
    crop_y = float(crop["y"])
    crop_w = float(crop["width"])
    crop_h = float(crop["height"])

    visible_source_w = max(1.0, source_w * crop_w)
    visible_source_h = max(1.0, source_h * crop_h)
    if image_mode == "裁剪填充":
        base_scale = max(
            width_cm / visible_source_w,
            height_cm / visible_source_h,
        )
    else:
        base_scale = min(
            width_cm / visible_source_w,
            height_cm / visible_source_h,
        )

    zoom = max(0.10, min(10.00, float(transform.get("zoom", 1.0))))
    offset_x = max(-50.0, min(50.0, float(transform.get("offset_x", 0.0))))
    offset_y = max(-50.0, min(50.0, float(transform.get("offset_y", 0.0))))
    scale = max(0.000001, base_scale * zoom)
    display_w = visible_source_w * scale
    display_h = visible_source_h * scale
    image_left = left_cm + (width_cm - display_w) / 2.0 + offset_x * width_cm * 0.5
    image_top = top_cm + (height_cm - display_h) / 2.0 + offset_y * height_cm * 0.5

    visible_left = max(left_cm, image_left)
    visible_top = max(top_cm, image_top)
    visible_right = min(left_cm + width_cm, image_left + display_w)
    visible_bottom = min(top_cm + height_cm, image_top + display_h)
    if visible_right <= visible_left or visible_bottom <= visible_top:
        # Extreme pan values can move the complete image outside its frame.
        # Keep the full source embedded as a tiny off-frame item so reset crop
        # remains possible without showing pixels that are absent in preview.
        return slide.shapes.add_picture(
            src_path,
            Cm(-1.0),
            Cm(-1.0),
            width=Cm(0.01),
            height=Cm(0.01),
        )

    within_left = (visible_left - image_left) / display_w
    within_top = (visible_top - image_top) / display_h
    within_right = (visible_right - image_left) / display_w
    within_bottom = (visible_bottom - image_top) / display_h
    source_left = crop_x + crop_w * within_left
    source_top = crop_y + crop_h * within_top
    source_right = crop_x + crop_w * within_right
    source_bottom = crop_y + crop_h * within_bottom

    picture = slide.shapes.add_picture(
        src_path,
        Cm(visible_left),
        Cm(visible_top),
        width=Cm(visible_right - visible_left),
        height=Cm(visible_bottom - visible_top),
    )
    picture.crop_left = max(0.0, min(1.0, source_left))
    picture.crop_top = max(0.0, min(1.0, source_top))
    picture.crop_right = max(0.0, min(1.0, 1.0 - source_right))
    picture.crop_bottom = max(0.0, min(1.0, 1.0 - source_bottom))
    return picture


def _add_compatibility_picture(
    slide,
    src_path: str,
    *,
    left_cm: float,
    top_cm: float,
    width_cm: float,
    height_cm: float,
    image_mode: str,
    transform: dict,
    export_dpi: int,
    temporary_files: list[str],
):
    """
    把MPO、WEBP、HEIF或python-pptx不支持的图片转为高清PNG，
    再写入PPT。

    只创建临时导出文件，不修改原始图片。
    """
    rendered = _render_image_box(
        src_path,
        width_cm,
        height_cm,
        image_mode,
        transform,
        export_dpi=export_dpi,
    )
    temporary_files.append(
        rendered
    )

    return slide.shapes.add_picture(
        rendered,
        Cm(left_cm),
        Cm(top_cm),
        width=Cm(width_cm),
        height=Cm(height_cm),
    )


def _add_original_picture(
    slide,
    src_path: str,
    *,
    left_cm: float,
    top_cm: float,
    width_cm: float,
    height_cm: float,
    image_mode: str,
):
    """
    未编辑图片直接嵌入原始文件，不再转成临时PNG。
    这样PPTX媒体包会保留原始像素和JPEG质量。
    """
    source_w, source_h = _image_pixel_size(
        src_path
    )
    source_ratio = source_w / source_h
    frame_ratio = width_cm / height_cm

    if image_mode == "裁剪填充":
        picture = slide.shapes.add_picture(
            src_path,
            Cm(left_cm),
            Cm(top_cm),
            width=Cm(width_cm),
            height=Cm(height_cm),
        )

        if source_ratio > frame_ratio:
            visible_fraction = (
                frame_ratio / source_ratio
            )
            crop = max(
                0.0,
                (1.0 - visible_fraction) / 2.0,
            )
            picture.crop_left = crop
            picture.crop_right = crop

        elif source_ratio < frame_ratio:
            visible_fraction = (
                source_ratio / frame_ratio
            )
            crop = max(
                0.0,
                (1.0 - visible_fraction) / 2.0,
            )
            picture.crop_top = crop
            picture.crop_bottom = crop

        return picture

    if source_ratio >= frame_ratio:
        picture_w = width_cm
        picture_h = width_cm / source_ratio
    else:
        picture_h = height_cm
        picture_w = height_cm * source_ratio

    picture_left = (
        left_cm
        + (width_cm - picture_w) / 2
    )
    picture_top = (
        top_cm
        + (height_cm - picture_h) / 2
    )

    return slide.shapes.add_picture(
        src_path,
        Cm(picture_left),
        Cm(picture_top),
        width=Cm(picture_w),
        height=Cm(picture_h),
    )


def _permission_error_message(
    output_path: Path,
    *,
    recovery_path: Path | None = None,
) -> str:
    message = (
        f"无法写入目标文件：\n{output_path}\n\n"
        "常见原因：\n"
        "1. 该PPT正被 WPS 或 PowerPoint 打开；\n"
        "2. 文件被设置为只读；\n"
        "3. 当前账号没有该文件夹的写入权限；\n"
        "4. 安全软件暂时锁定了该文件。\n\n"
        "请先关闭正在打开的PPT，或更换输出文件名后重新生成。"
    )

    if recovery_path is not None:
        message += (
            "\n\n目标文件在最终写入时被占用，"
            "本次已经生成完成的完整PPT已保存在：\n"
            f"{recovery_path}"
        )

    return message


def preflight_ppt_output(
    output_file: str | os.PathLike,
) -> str:
    """
    在开始生成PPT前验证目标目录和目标文件是否可写。

    该检查不会删除或覆盖现有文件。
    """
    output_path = Path(output_file)

    if not str(output_path).strip():
        raise ValueError(
            "尚未设置PPT输出文件。"
        )

    if output_path.exists() and output_path.is_dir():
        raise PermissionError(
            f"输出路径是文件夹，不是PPT文件：\n{output_path}"
        )

    parent = output_path.parent
    parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # 验证目录可创建和删除临时文件。
    probe_path = None

    try:
        file_descriptor, probe_name = (
            tempfile.mkstemp(
                prefix=".framedeck_write_probe_",
                suffix=".tmp",
                dir=str(parent),
            )
        )
        os.close(file_descriptor)
        probe_path = Path(probe_name)

    except PermissionError as error:
        raise PermissionError(
            _permission_error_message(
                output_path
            )
        ) from error

    finally:
        if probe_path is not None:
            try:
                probe_path.unlink(
                    missing_ok=True
                )
            except OSError:
                pass

    # 目标文件已经存在时，检测它是否正被独占锁定。
    if output_path.exists():
        file_descriptor = None

        try:
            flags = os.O_RDWR

            if hasattr(os, "O_BINARY"):
                flags |= os.O_BINARY

            file_descriptor = os.open(
                str(output_path),
                flags,
            )

        except PermissionError as error:
            raise PermissionError(
                _permission_error_message(
                    output_path
                )
            ) from error

        finally:
            if file_descriptor is not None:
                os.close(
                    file_descriptor
                )

    return str(output_path)


def _atomic_save_presentation(
    presentation,
    output_path: Path,
) -> str:
    """
    先写入同目录临时PPT，再原子替换目标文件。

    优点：
    - 保存中断不会留下半个损坏的 output.pptx；
    - 目标文件在生成期间突然被打开时，将完整结果保留为恢复文件。
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary_path = output_path.with_name(
        (
            f".{output_path.stem}."
            f"framedeck_{os.getpid()}_{time.time_ns()}"
            f"{output_path.suffix or '.pptx'}"
        )
    )

    recovery_path = None

    try:
        presentation.save(
            str(temporary_path)
        )

        try:
            os.replace(
                temporary_path,
                output_path,
            )

        except PermissionError as error:
            timestamp = time.strftime(
                "%Y%m%d_%H%M%S"
            )
            recovery_path = output_path.with_name(
                (
                    f"{output_path.stem}"
                    f"_FrameDeck_恢复_{timestamp}"
                    f"{output_path.suffix or '.pptx'}"
                )
            )

            # 避免极低概率的同名冲突。
            counter = 2

            while recovery_path.exists():
                recovery_path = output_path.with_name(
                    (
                        f"{output_path.stem}"
                        f"_FrameDeck_恢复_{timestamp}_{counter}"
                        f"{output_path.suffix or '.pptx'}"
                    )
                )
                counter += 1

            os.replace(
                temporary_path,
                recovery_path,
            )

            raise PermissionError(
                _permission_error_message(
                    output_path,
                    recovery_path=recovery_path,
                )
            ) from error

        return str(output_path)

    finally:
        if temporary_path.exists():
            try:
                temporary_path.unlink(
                    missing_ok=True
                )
            except OSError:
                pass


def _page_ranges_for_export(
    image_count: int,
    per_page: int,
    page_breaks=None,
):
    image_count = max(0, int(image_count))
    per_page = max(1, int(per_page))

    if image_count <= 0:
        return []

    boundaries = [0]
    boundaries.extend(
        sorted(
            {
                int(value)
                for value in list(page_breaks or [])
                if 0 < int(value) < image_count
            }
        )
    )
    boundaries.append(image_count)

    ranges = []

    for start, end in zip(
        boundaries[:-1],
        boundaries[1:],
    ):
        cursor = start

        while cursor < end:
            page_end = min(
                cursor + per_page,
                end,
            )
            ranges.append(
                (cursor, page_end)
            )
            cursor = page_end

    return ranges


def _logical_export_page_plan(
    page_ranges,
    blank_page_positions=None,
):
    """
    将真实图片页与空白页合并成最终逻辑页面顺序。

    返回：
        [
            ("images", (0, 10)),
            ("blank", None),
            ("images", (10, 15)),
        ]
    """
    page_ranges = list(
        page_ranges
        or []
    )
    base_count = len(
        page_ranges
    )

    raw_positions = sorted(
        {
            int(value)
            for value in list(
                blank_page_positions
                or []
            )
            if int(value) >= 0
        }
    )
    normalized = []

    for value in raw_positions:
        maximum = (
            base_count
            + len(
                normalized
            )
        )
        value = max(
            0,
            min(
                value,
                maximum,
            ),
        )

        if value in normalized:
            continue

        normalized.append(
            value
        )

    blank_set = set(
        normalized
    )
    total_count = (
        base_count
        + len(
            normalized
        )
    )
    plan = []
    image_cursor = 0

    for logical_index in range(
        total_count
    ):
        if logical_index in blank_set:
            plan.append(
                (
                    "blank",
                    None,
                )
            )
        elif image_cursor < base_count:
            plan.append(
                (
                    "images",
                    page_ranges[
                        image_cursor
                    ],
                )
            )
            image_cursor += 1

    return plan


def generate_ppt(
    image_paths: list[str],
    output_file: str,
    image_transforms: dict[str, dict] | None = None,
    hidden_images: list[str] | None = None,
    show_title: bool = False,
    title_text: str = "",
    subtitle_text: str = "",
    title_height: float = 0.9,
    show_footer: bool = False,
    footer_text: str = "",
    logo_path: str = "",
    logo_width: float = 1.5,
    rows: int = 2,
    cols: int = 6,
    page_size: str = "16:9",
    image_mode: str = "保持比例",
    show_filename: bool = False,
    show_page_number: bool = True,
    show_index: bool = False,
    add_border: bool = False,
    margin_left: float = 0.6,
    margin_right: float = 0.6,
    margin_top: float = 0.8,
    margin_bottom: float = 0.8,
    gap_x: float = 0.25,
    gap_y: float = 0.35,
    label_height: float = 0.45,
    progress_callback: Callable[[int], None] | None = None,
    log_callback: Callable[[str], None] | None = None,
    show_grid: bool = False,
    ppt_quality_dpi: int = PPT_EXPORT_DPI,
    preserve_original_images: bool = True,
    page_breaks: list[int] | None = None,
    resolved_title_pages: list[dict] | None = None,
    title_system: dict | None = None,
    image_groups: list[dict] | None = None,
    blank_page_positions: list[int] | None = None,
    **kwargs,
) -> str:
    # UI-05-45A：页面 Logo 与副标题已从产品界面移除。
    # 这里主动忽略旧工程或旧调用方残留的字段，确保预览与导出一致。
    subtitle_text = ""
    logo_path = ""

    output_path = Path(
        preflight_ppt_output(
            output_file
        )
    )

    hidden = set(hidden_images or [])
    images = [
        p
        for p in image_paths
        if Path(p).is_file()
        and p not in hidden
    ]

    if log_callback:
        log_callback(
            f"输出路径检查通过：{output_path}"
        )

    if not images:
        raise ValueError("没有可用于生成 PPT 的可见图片。")
    if rows < 1 or cols < 1:
        raise ValueError("行数和列数必须大于 0。")

    transforms = image_transforms or {}
    slide_w, slide_h = PAGE_SIZES.get(
        page_size,
        PAGE_SIZES["16:9"],
    )

    usable_w = (
        slide_w
        - margin_left
        - margin_right
        - (cols - 1) * gap_x
    )

    if usable_w <= 0:
        raise ValueError(
            "页边距或水平间距过大，已没有可用排版区域。"
        )

    cell_w = usable_w / cols
    use_label = (
        show_filename
        or show_index
    )

    def page_title_context(
        page_index: int,
    ) -> dict:
        contexts = list(
            resolved_title_pages
            or []
        )

        if (
            0
            <= page_index
            < len(contexts)
            and isinstance(
                contexts[
                    page_index
                ],
                dict,
            )
        ):
            context = dict(
                contexts[
                    page_index
                ]
            )
            context["subtitle"] = ""
            return context

        return {
            "title": (
                title_text.strip()
                if show_title
                else ""
            ),
            "subtitle": "",
            "style": {
                "font_family": (
                    "Microsoft YaHei"
                ),
                "font_size": 20.0,
                "color": "#111827",
                "bold": True,
                "alignment": "center",
                "top_spacing_cm": 0.0,
                "region_height_cm": (
                    float(
                        title_height
                    )
                ),
            },
        }

    def page_geometry(
        page_index: int,
    ) -> dict:
        context = page_title_context(
            page_index
        )
        style = dict(
            context.get(
                "style",
                {},
            )
            or {}
        )
        page_title = str(
            context.get(
                "title",
                "",
            )
            or ""
        ).strip()
        page_subtitle = ""

        try:
            region_height = float(
                style.get(
                    "region_height_cm",
                    title_height,
                )
            )
        except (TypeError, ValueError):
            region_height = float(
                title_height
            )

        try:
            top_spacing = float(
                style.get(
                    "top_spacing_cm",
                    0.0,
                )
            )
        except (TypeError, ValueError):
            top_spacing = 0.0

        logo_exists = False

        title_geometry = resolve_title_vertical_geometry(
            {
                **style,
                "region_height_cm": region_height,
                "top_spacing_cm": top_spacing,
            },
            enabled=show_title,
            has_content=bool(page_title or page_subtitle or logo_exists),
        )
        effective_title_h = float(title_geometry["effective_height_cm"])
        effective_top_spacing = float(
            title_geometry["effective_top_spacing_cm"]
        )

        usable_h = (
            slide_h
            - margin_top
            - margin_bottom
            - effective_top_spacing
            - effective_title_h
            - (rows - 1) * gap_y
        )

        if usable_h <= 0:
            raise ValueError(
                (
                    f"第 {page_index + 1} 页的标题高度、"
                    "页边距或间距过大，已没有可用排版区域。"
                )
            )

        cell_h = (
            usable_h
            / rows
        )
        actual_label_h = (
            min(
                label_height,
                max(
                    0.0,
                    cell_h * 0.30,
                ),
            )
            if use_label
            else 0.0
        )
        image_area_h = (
            cell_h
            - actual_label_h
        )

        if image_area_h <= 0:
            raise ValueError(
                f"第 {page_index + 1} 页文字区域过高，图片区域不足。"
            )

        return {
            "title": page_title,
            "subtitle": page_subtitle,
            "style": style,
            "effective_title_h": (
                effective_title_h
            ),
            "effective_top_spacing": (
                effective_top_spacing
            ),
            "cell_h": cell_h,
            "actual_label_h": (
                actual_label_h
            ),
            "image_area_h": (
                image_area_h
            ),
        }

    prs = Presentation()
    prs.slide_width = Cm(slide_w)
    prs.slide_height = Cm(slide_h)

    per_page = rows * cols
    page_ranges = _page_ranges_for_export(
        len(images),
        per_page,
        page_breaks,
    )
    logical_page_plan = (
        _logical_export_page_plan(
            page_ranges,
            blank_page_positions,
        )
    )
    total_pages = len(
        logical_page_plan
    )

    placements = []

    for page_no, (
        page_kind,
        page_range_value,
    ) in enumerate(
        logical_page_plan,
        start=1,
    ):
        if page_kind == "blank":
            placements.append(
                (
                    None,
                    None,
                    page_no,
                    -1,
                    True,
                    True,
                )
            )
            continue

        page_start, page_end = (
            page_range_value
        )

        for position, idx in enumerate(
            range(
                page_start,
                page_end,
            )
        ):
            placements.append(
                (
                    idx,
                    images[idx],
                    page_no,
                    position,
                    position == 0,
                    False,
                )
            )

    temporary_files: list[str] = []

    try:
        slide = None
        for (
            idx,
            image_path,
            page_no,
            position,
            starts_page,
            is_blank_page,
        ) in placements:
            if starts_page:
                slide = prs.slides.add_slide(
                    prs.slide_layouts[6]
                )
                current_geometry = page_geometry(
                    page_no - 1
                )

                page_title = str(
                    current_geometry[
                        "title"
                    ]
                    or ""
                ).strip()
                page_subtitle = str(
                    current_geometry[
                        "subtitle"
                    ]
                    or ""
                ).strip()
                title_style = dict(
                    current_geometry[
                        "style"
                    ]
                    or {}
                )

                if (
                    not is_blank_page
                    and show_title
                    and (
                        page_title
                        or page_subtitle
                        or (
                            logo_path
                            and Path(
                                logo_path
                            ).is_file()
                        )
                    )
                ):
                    logo_space = 0.0

                    if (
                        logo_path
                        and Path(
                            logo_path
                        ).is_file()
                    ):
                        logo_space = (
                            logo_width
                            + 0.25
                        )
                        slide.shapes.add_picture(
                            logo_path,
                            Cm(margin_left),
                            Cm(
                                margin_top
                                + current_geometry[
                                    "effective_top_spacing"
                                ]
                            ),
                            width=Cm(
                                logo_width
                            ),
                        )

                    title_box = slide.shapes.add_textbox(
                        Cm(
                            margin_left
                            + logo_space
                        ),
                        Cm(
                            margin_top
                            + current_geometry[
                                "effective_top_spacing"
                            ]
                        ),
                        Cm(
                            slide_w
                            - margin_left
                            - margin_right
                            - logo_space
                        ),
                        Cm(
                            current_geometry[
                                "effective_title_h"
                            ]
                        ),
                    )
                    title_frame = (
                        title_box.text_frame
                    )
                    title_frame.clear()
                    title_frame.margin_left = 0
                    title_frame.margin_right = 0
                    title_frame.margin_top = 0
                    title_frame.margin_bottom = 0
                    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    title_frame.word_wrap = False
                    title_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

                    alignment = str(
                        title_style.get(
                            "alignment",
                            "center",
                        )
                        or "center"
                    ).lower()
                    ppt_alignment = {
                        "left": PP_ALIGN.LEFT,
                        "right": PP_ALIGN.RIGHT,
                    }.get(
                        alignment,
                        PP_ALIGN.CENTER,
                    )

                    font_family = str(
                        title_style.get(
                            "font_family",
                            "Microsoft YaHei",
                        )
                        or "Microsoft YaHei"
                    )
                    try:
                        font_size = float(
                            title_style.get(
                                "font_size",
                                20.0,
                            )
                        )
                    except (
                        TypeError,
                        ValueError,
                    ):
                        font_size = 20.0

                    color_text = str(
                        title_style.get(
                            "color",
                            "#111827",
                        )
                        or "#111827"
                    ).lstrip("#")
                    if not re.fullmatch(
                        r"[0-9A-Fa-f]{6}",
                        color_text,
                    ):
                        color_text = "111827"

                    title_paragraph = (
                        title_frame.paragraphs[
                            0
                        ]
                    )
                    title_paragraph.text = (
                        page_title
                    )
                    title_paragraph.alignment = (
                        ppt_alignment
                    )
                    title_paragraph.font.name = (
                        font_family
                    )
                    title_paragraph.font.size = Pt(
                        max(
                            6.0,
                            min(
                                96.0,
                                font_size,
                            ),
                        )
                    )
                    title_paragraph.font.bold = bool(
                        title_style.get(
                            "bold",
                            True,
                        )
                    )
                    title_paragraph.font.color.rgb = RGBColor.from_string(
                        color_text.upper()
                    )

                    if page_subtitle:
                        subtitle_paragraph = (
                            title_frame.add_paragraph()
                        )
                        subtitle_paragraph.text = (
                            page_subtitle
                        )
                        subtitle_paragraph.alignment = (
                            ppt_alignment
                        )
                        subtitle_paragraph.font.name = (
                            font_family
                        )
                        subtitle_paragraph.font.size = Pt(
                            max(
                                6.0,
                                font_size
                                * 0.55,
                            )
                        )
                        subtitle_paragraph.font.bold = (
                            False
                        )
                        subtitle_paragraph.font.color.rgb = RGBColor.from_string(
                            color_text.upper()
                        )

                if log_callback:
                    log_callback(
                        f"正在生成第 {page_no}/{total_pages} 页"
                    )

                if is_blank_page:
                    continue

            row = position // cols
            col = position % cols

            cell_h = current_geometry[
                "cell_h"
            ]
            actual_label_h = current_geometry[
                "actual_label_h"
            ]
            image_area_h = current_geometry[
                "image_area_h"
            ]

            cell_left = (
                margin_left
                + col
                * (
                    cell_w
                    + gap_x
                )
            )
            cell_top = (
                margin_top
                + current_geometry[
                    "effective_top_spacing"
                ]
                + current_geometry[
                    "effective_title_h"
                ]
                + row
                * (
                    cell_h
                    + gap_y
                )
            )

            transform = transforms.get(
                image_path,
                {},
            )

            requires_render = _requires_render(
                image_path,
                transform,
            )
            source_format = _source_image_format(
                str(image_path)
            )

            if (
                source_format
                and source_format
                not in PPT_DIRECT_IMAGE_FORMATS
                and log_callback
            ):
                log_callback(
                    (
                        "兼容转换："
                        f"{Path(image_path).name} "
                        f"({source_format} → PNG)"
                    )
                )

            editable_source = str(image_path)
            editable_transform = dict(transform)

            if requires_render:
                editable_source = _render_full_editable_source(
                    str(image_path),
                    transform,
                )
                temporary_files.append(editable_source)
                editable_transform["crop"] = _oriented_crop(
                    transform.get("crop"),
                    rotation=int(transform.get("rotation", 0)),
                    flip_h=bool(transform.get("flip_h", False)),
                    flip_v=bool(transform.get("flip_v", False)),
                )
                editable_transform["rotation"] = 0
                editable_transform["flip_h"] = False
                editable_transform["flip_v"] = False

            try:
                picture = _add_editable_picture(
                    slide,
                    editable_source,
                    left_cm=cell_left,
                    top_cm=cell_top,
                    width_cm=cell_w,
                    height_cm=image_area_h,
                    image_mode=image_mode,
                    transform=editable_transform,
                )
            except (
                ValueError,
                OSError,
                UnidentifiedImageError,
            ) as direct_error:
                if editable_source != str(image_path):
                    raise
                if log_callback:
                    log_callback(
                        (
                            "原图直写不兼容，完整转换后保留可编辑裁切："
                            f"{Path(image_path).name} · {direct_error}"
                        )
                    )
                editable_source = _render_full_editable_source(
                    str(image_path),
                    transform,
                )
                temporary_files.append(editable_source)
                editable_transform["crop"] = _oriented_crop(
                    transform.get("crop"),
                    rotation=int(transform.get("rotation", 0)),
                    flip_h=bool(transform.get("flip_h", False)),
                    flip_v=bool(transform.get("flip_v", False)),
                )
                picture = _add_editable_picture(
                    slide,
                    editable_source,
                    left_cm=cell_left,
                    top_cm=cell_top,
                    width_cm=cell_w,
                    height_cm=image_area_h,
                    image_mode=image_mode,
                    transform=editable_transform,
                )

            if add_border:
                picture.line.width = Pt(0.8)

            if use_label:
                parts = []
                if show_index:
                    parts.append(f"{idx + 1:03d}")
                if show_filename:
                    parts.append(Path(image_path).stem)

                textbox = slide.shapes.add_textbox(
                    Cm(cell_left),
                    Cm(cell_top + image_area_h),
                    Cm(cell_w),
                    Cm(actual_label_h),
                )
                frame = textbox.text_frame
                frame.clear()
                frame.margin_left = 0
                frame.margin_right = 0
                frame.margin_top = 0
                frame.margin_bottom = 0
                paragraph = frame.paragraphs[0]
                paragraph.text = "  ".join(parts)
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.size = Pt(8)

            if progress_callback:
                progress_callback(round((idx + 1) / len(images) * 100))

        if show_footer and footer_text.strip():
            for current_slide in prs.slides:
                footer_box = current_slide.shapes.add_textbox(
                    Cm(margin_left),
                    Cm(slide_h - margin_bottom + 0.05),
                    Cm(slide_w - margin_left - margin_right),
                    Cm(0.28),
                )
                footer_frame = footer_box.text_frame
                footer_frame.clear()
                footer_frame.margin_left = 0
                footer_frame.margin_right = 0
                footer_frame.margin_top = 0
                footer_frame.margin_bottom = 0
                footer_p = footer_frame.paragraphs[0]
                footer_p.text = footer_text.strip()
                footer_p.alignment = PP_ALIGN.LEFT
                footer_p.font.size = Pt(7)

        if show_page_number:
            for page_index, current_slide in enumerate(prs.slides, start=1):
                textbox = current_slide.shapes.add_textbox(
                    Cm(slide_w / 2 - 1.5),
                    Cm(slide_h - 0.48),
                    Cm(3),
                    Cm(0.3),
                )
                frame = textbox.text_frame
                frame.clear()
                frame.margin_left = 0
                frame.margin_right = 0
                frame.margin_top = 0
                frame.margin_bottom = 0
                paragraph = frame.paragraphs[0]
                paragraph.text = f"{page_index} / {total_pages}"
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.size = Pt(8)

        saved_output = (
            _atomic_save_presentation(
                prs,
                output_path,
            )
        )

        if log_callback:
            log_callback(
                f"完成：{saved_output}"
            )
        return str(saved_output)

    finally:
        for temporary_file in temporary_files:
            try:
                os.remove(temporary_file)
            except OSError:
                pass
