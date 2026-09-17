"""
FrameDeck Studio
UI-05-41A - PPT 图片素材导入核心

安装位置：
    core/ppt_importer.py

功能：
- 只读打开 .pptx / .pptm；
- 按页码和页面视觉位置提取图片；
- 每个图片对象生成独立文件，重复使用的图片也可分别导入；
- 尽量保留 PPT 中的裁切、旋转与镜像参数；
- 不执行宏，不修改原始 PPT；
- 提取文件保存在 FrameDeck 用户数据目录，工程关闭后仍可继续使用。
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import time
import traceback
import uuid
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from statistics import median
from typing import Any, Callable, Iterable

from PIL import (
    Image,
    ImageOps,
    UnidentifiedImageError,
)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.app_paths import user_data_dir


PPT_IMPORT_MODE_FLAT = "flat_group"
PPT_IMPORT_MODE_GROUP_BY_SLIDE = "group_by_slide"
PPT_IMPORT_MODE_PAGE_BY_SLIDE = "page_by_slide"

VALID_PPT_IMPORT_MODES = frozenset(
    {
        PPT_IMPORT_MODE_FLAT,
        PPT_IMPORT_MODE_GROUP_BY_SLIDE,
        PPT_IMPORT_MODE_PAGE_BY_SLIDE,
    }
)


SUPPORTED_PPT_EXTENSIONS = frozenset(
    {
        ".pptx",
        ".pptm",
    }
)

RAW_PRESERVE_FORMATS = frozenset(
    {
        "JPEG",
        "PNG",
        "BMP",
        "TIFF",
        "WEBP",
    }
)

FORMAT_EXTENSION = {
    "JPEG": ".jpg",
    "PNG": ".png",
    "BMP": ".bmp",
    "TIFF": ".tif",
    "WEBP": ".webp",
}


class PPTImportCancelled(RuntimeError):
    """用户主动取消PPT图片提取。"""


@dataclass
class PictureRecord:
    slide_index: int
    shape_index: int
    shape_name: str
    left: int
    top: int
    width: int
    height: int
    rotation: float
    flip_h: bool
    flip_v: bool
    crop_left: float
    crop_top: float
    crop_right: float
    crop_bottom: float
    blob: bytes
    relationship_id: str
    visual_order: int = 0

    @property
    def center_y(self) -> float:
        return float(self.top) + float(self.height) / 2.0


def _safe_name(
    value: str,
    *,
    fallback: str = "Picture",
    max_length: int = 48,
) -> str:
    value = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]+',
        "_",
        str(value or ""),
    )
    value = re.sub(
        r"\s+",
        " ",
        value,
    ).strip(" ._")

    if not value:
        value = fallback

    return value[:max_length]


def _clamp(
    value: Any,
    minimum: float = 0.0,
    maximum: float = 1.0,
) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        number = 0.0

    return max(
        minimum,
        min(
            maximum,
            number,
        ),
    )


def _shape_flip_state(
    shape,
) -> tuple[bool, bool]:
    try:
        transform = shape._element.spPr.xfrm
    except Exception:
        return False, False

    if transform is None:
        return False, False

    def truthy(attribute: str) -> bool:
        value = str(
            transform.get(attribute) or ""
        ).strip().lower()
        return value in {
            "1",
            "true",
            "yes",
        }

    return (
        truthy("flipH"),
        truthy("flipV"),
    )


def _iter_picture_shapes(
    shapes,
    *,
    parent_left: int = 0,
    parent_top: int = 0,
    index_prefix: int = 0,
) -> Iterable[tuple[int, Any, int, int]]:
    """
    递归查找普通图片和组合对象中的图片。

    组合对象内部坐标在不同PPT中可能略有差异；
    这里只使用近似视觉坐标进行排序，不影响图片数据提取。
    """
    for local_index, shape in enumerate(
        shapes
    ):
        shape_index = (
            index_prefix * 10000
            + local_index
        )

        try:
            shape_type = shape.shape_type
        except Exception:
            continue

        if (
            shape_type
            == MSO_SHAPE_TYPE.PICTURE
        ):
            yield (
                shape_index,
                shape,
                parent_left
                + int(getattr(shape, "left", 0)),
                parent_top
                + int(getattr(shape, "top", 0)),
            )
            continue

        if (
            shape_type
            == MSO_SHAPE_TYPE.GROUP
        ):
            try:
                group_shapes = shape.shapes
            except Exception:
                continue

            yield from _iter_picture_shapes(
                group_shapes,
                parent_left=(
                    parent_left
                    + int(
                        getattr(
                            shape,
                            "left",
                            0,
                        )
                    )
                ),
                parent_top=(
                    parent_top
                    + int(
                        getattr(
                            shape,
                            "top",
                            0,
                        )
                    )
                ),
                index_prefix=shape_index + 1,
            )


def _visual_sort(
    records: list[PictureRecord],
) -> list[PictureRecord]:
    """
    按视觉阅读顺序排序：
    - 先按大致行分组；
    - 行内从左到右；
    - 行之间从上到下。

    相比简单 top/left 排序，对错落排版更稳定。
    """
    if len(records) <= 1:
        return list(records)

    heights = [
        max(1, int(record.height))
        for record in records
    ]
    reference_height = float(
        median(heights)
    )
    base_tolerance = max(
        1.0,
        reference_height * 0.36,
    )

    rows: list[dict[str, Any]] = []

    for record in sorted(
        records,
        key=lambda item: (
            item.center_y,
            item.left,
            item.shape_index,
        ),
    ):
        best_row = None
        best_distance = None

        for row in rows:
            tolerance = max(
                base_tolerance,
                float(record.height) * 0.30,
                float(row["height"]) * 0.30,
            )
            distance = abs(
                record.center_y
                - float(row["center_y"])
            )

            if (
                distance <= tolerance
                and (
                    best_distance is None
                    or distance < best_distance
                )
            ):
                best_row = row
                best_distance = distance

        if best_row is None:
            rows.append(
                {
                    "center_y": record.center_y,
                    "height": float(
                        record.height
                    ),
                    "items": [record],
                }
            )
            continue

        best_row["items"].append(
            record
        )
        count = len(
            best_row["items"]
        )
        best_row["center_y"] = (
            (
                float(
                    best_row["center_y"]
                )
                * (count - 1)
                + record.center_y
            )
            / count
        )
        best_row["height"] = max(
            float(best_row["height"]),
            float(record.height),
        )

    rows.sort(
        key=lambda row: float(
            row["center_y"]
        )
    )

    ordered: list[PictureRecord] = []

    for row in rows:
        row_items = sorted(
            row["items"],
            key=lambda item: (
                item.left,
                item.top,
                item.shape_index,
            ),
        )
        ordered.extend(
            row_items
        )

    for visual_order, record in enumerate(
        ordered,
        start=1,
    ):
        record.visual_order = (
            visual_order
        )

    return ordered


def _collect_picture_records(
    presentation,
    *,
    progress_callback: (
        Callable[[int, int, str], None]
        | None
    ) = None,
    cancel_callback: (
        Callable[[], bool]
        | None
    ) = None,
) -> tuple[
    list[PictureRecord],
    list[str],
]:
    warnings: list[str] = []
    records: list[PictureRecord] = []
    slide_total = len(
        presentation.slides
    )

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        if (
            cancel_callback
            and cancel_callback()
        ):
            raise PPTImportCancelled(
                "用户取消PPT图片导入。"
            )

        if progress_callback:
            progress_callback(
                slide_number - 1,
                max(1, slide_total),
                (
                    "正在扫描PPT页面："
                    f"{slide_number}/{slide_total}"
                ),
            )

        slide_records: list[
            PictureRecord
        ] = []

        for (
            shape_index,
            shape,
            effective_left,
            effective_top,
        ) in _iter_picture_shapes(
            slide.shapes
        ):
            try:
                image = shape.image
                blob = bytes(image.blob)
            except Exception as error:
                warnings.append(
                    (
                        f"第{slide_number}页的图片对象"
                        f"“{getattr(shape, 'name', '')}”读取失败："
                        f"{type(error).__name__}: {error}"
                    )
                )
                continue

            flip_h, flip_v = (
                _shape_flip_state(
                    shape
                )
            )

            try:
                relationship_id = str(
                    shape._element.blipFill.blip.embed
                    or ""
                )
            except Exception:
                relationship_id = ""

            slide_records.append(
                PictureRecord(
                    slide_index=slide_number,
                    shape_index=shape_index,
                    shape_name=str(
                        getattr(
                            shape,
                            "name",
                            "",
                        )
                        or ""
                    ),
                    left=effective_left,
                    top=effective_top,
                    width=max(
                        1,
                        int(
                            getattr(
                                shape,
                                "width",
                                1,
                            )
                        ),
                    ),
                    height=max(
                        1,
                        int(
                            getattr(
                                shape,
                                "height",
                                1,
                            )
                        ),
                    ),
                    rotation=float(
                        getattr(
                            shape,
                            "rotation",
                            0.0,
                        )
                        or 0.0
                    ),
                    flip_h=flip_h,
                    flip_v=flip_v,
                    crop_left=float(
                        getattr(
                            shape,
                            "crop_left",
                            0.0,
                        )
                        or 0.0
                    ),
                    crop_top=float(
                        getattr(
                            shape,
                            "crop_top",
                            0.0,
                        )
                        or 0.0
                    ),
                    crop_right=float(
                        getattr(
                            shape,
                            "crop_right",
                            0.0,
                        )
                        or 0.0
                    ),
                    crop_bottom=float(
                        getattr(
                            shape,
                            "crop_bottom",
                            0.0,
                        )
                        or 0.0
                    ),
                    blob=blob,
                    relationship_id=relationship_id,
                )
            )

        records.extend(
            _visual_sort(
                slide_records
            )
        )

    if progress_callback:
        progress_callback(
            slide_total,
            max(1, slide_total),
            (
                "PPT扫描完成，"
                f"发现 {len(records)} 个图片对象"
            ),
        )

    return records, warnings


def build_slide_summaries(
    records: Iterable[PictureRecord],
    extracted: Iterable[dict[str, Any]],
    slide_count: int,
) -> list[dict[str, Any]]:
    """
    UI-05-43A：
    为 FrameDeck 返回完整的原 PPT 页面结构。

    即使某一页没有图片，也保留该页记录，便于
    “原PPT每页直接对应FrameDeck页面”模式生成真正空白页。
    """
    slide_count = max(
        0,
        int(slide_count),
    )

    object_counts = {
        slide_index: 0
        for slide_index in range(
            1,
            slide_count + 1,
        )
    }
    extracted_by_slide = {
        slide_index: []
        for slide_index in range(
            1,
            slide_count + 1,
        )
    }

    for record in list(
        records or []
    ):
        slide_index = int(
            getattr(
                record,
                "slide_index",
                0,
            )
            or 0
        )

        if 1 <= slide_index <= slide_count:
            object_counts[
                slide_index
            ] += 1

    for item in list(
        extracted or []
    ):
        if not isinstance(
            item,
            dict,
        ):
            continue

        try:
            slide_index = int(
                item.get(
                    "slide_index",
                    0,
                )
                or 0
            )
        except (
            TypeError,
            ValueError,
        ):
            continue

        if not (
            1
            <= slide_index
            <= slide_count
        ):
            continue

        extracted_by_slide[
            slide_index
        ].append(
            item
        )

    summaries = []

    for slide_index in range(
        1,
        slide_count + 1,
    ):
        items = sorted(
            extracted_by_slide[
                slide_index
            ],
            key=lambda item: (
                int(
                    item.get(
                        "visual_order",
                        0,
                    )
                    or 0
                ),
                str(
                    item.get(
                        "path",
                        "",
                    )
                ),
            ),
        )

        paths = [
            str(
                item.get(
                    "path",
                    "",
                )
            )
            for item in items
            if item.get(
                "path"
            )
        ]

        object_count = int(
            object_counts.get(
                slide_index,
                0,
            )
        )
        extracted_count = len(
            paths
        )

        summaries.append(
            {
                "slide_index": (
                    slide_index
                ),
                "picture_object_count": (
                    object_count
                ),
                "extracted_count": (
                    extracted_count
                ),
                "skipped_count": max(
                    0,
                    object_count
                    - extracted_count,
                ),
                "paths": paths,
                "visual_orders": [
                    int(
                        item.get(
                            "visual_order",
                            0,
                        )
                        or 0
                    )
                    for item in items
                ],
            }
        )

    return summaries


def _blob_format(
    blob: bytes,
) -> tuple[str, str]:
    """
    返回：
        (真实格式, 建议扩展名)
    """
    try:
        with Image.open(
            BytesIO(blob)
        ) as image:
            image_format = str(
                image.format or ""
            ).upper()

        return (
            image_format,
            FORMAT_EXTENSION.get(
                image_format,
                "",
            ),
        )

    except Exception:
        return "", ""


def _write_picture_blob(
    blob: bytes,
    base_path: Path,
) -> tuple[
    Path | None,
    str,
    str | None,
]:
    """
    保存一个PPT图片对象。

    返回：
        (文件路径, 真实格式, 警告)
    """
    image_format, raw_extension = (
        _blob_format(blob)
    )

    if (
        image_format
        in RAW_PRESERVE_FORMATS
        and raw_extension
    ):
        output_path = (
            base_path.with_suffix(
                raw_extension
            )
        )
        output_path.write_bytes(
            blob
        )
        return (
            output_path,
            image_format,
            None,
        )

    # MPO、GIF及其他Pillow可读格式统一转为无损PNG。
    try:
        with Image.open(
            BytesIO(blob)
        ) as opened:
            try:
                opened.seek(0)
            except (
                EOFError,
                AttributeError,
            ):
                pass

            opened.load()
            converted = (
                ImageOps.exif_transpose(
                    opened
                )
            )

            if converted.mode not in {
                "RGB",
                "RGBA",
                "L",
                "LA",
            }:
                converted = converted.convert(
                    "RGBA"
                    if "A" in converted.getbands()
                    else "RGB"
                )

            output_path = (
                base_path.with_suffix(
                    ".png"
                )
            )
            converted.save(
                output_path,
                "PNG",
                optimize=False,
            )

        return (
            output_path,
            image_format or "UNKNOWN",
            (
                None
                if image_format
                else "图片格式未标记，已转换为PNG。"
            ),
        )

    except (
        UnidentifiedImageError,
        OSError,
        ValueError,
    ) as error:
        return (
            None,
            image_format or "UNKNOWN",
            (
                "该图片对象可能是SVG、EMF、WMF或损坏媒体，"
                "当前版本未提取："
                f"{type(error).__name__}: {error}"
            ),
        )


def _transform_from_record(
    record: PictureRecord,
) -> dict[str, Any]:
    left = _clamp(
        record.crop_left
    )
    top = _clamp(
        record.crop_top
    )
    right = _clamp(
        record.crop_right
    )
    bottom = _clamp(
        record.crop_bottom
    )

    width = max(
        0.03,
        1.0 - left - right,
    )
    height = max(
        0.03,
        1.0 - top - bottom,
    )

    if left + width > 1.0:
        left = max(
            0.0,
            1.0 - width,
        )

    if top + height > 1.0:
        top = max(
            0.0,
            1.0 - height,
        )

    crop_modified = any(
        abs(value) > 1e-6
        for value in (
            left,
            top,
            right,
            bottom,
        )
    )

    return {
        "zoom": 1.0,
        "offset_x": 0.0,
        "offset_y": 0.0,
        "rotation": int(
            round(record.rotation)
        ) % 360,
        "flip_h": bool(
            record.flip_h
        ),
        "flip_v": bool(
            record.flip_v
        ),
        "crop": {
            "enabled": crop_modified,
            "x": round(left, 6),
            "y": round(top, 6),
            "width": round(width, 6),
            "height": round(height, 6),
            "aspect_mode": "free",
            "aspect_ratio": None,
        },
    }


def _import_directory_for(
    source_path: Path,
) -> Path:
    try:
        stat = source_path.stat()
        fingerprint_source = (
            f"{source_path.resolve()}|"
            f"{stat.st_size}|"
            f"{stat.st_mtime_ns}"
        )
    except OSError:
        fingerprint_source = str(
            source_path
        )

    fingerprint = hashlib.sha1(
        fingerprint_source.encode(
            "utf-8",
            errors="ignore",
        )
    ).hexdigest()[:10]

    stamp = time.strftime(
        "%Y%m%d_%H%M%S"
    )
    session_id = uuid.uuid4().hex[:8]
    folder_name = (
        f"{_safe_name(source_path.stem)}"
        f"_{stamp}_{fingerprint}_{session_id}"
    )

    output_dir = (
        user_data_dir()
        / "imports"
        / "ppt_images"
        / folder_name
    )
    output_dir.mkdir(
        parents=True,
        exist_ok=False,
    )
    return output_dir


def extract_ppt_images(
    ppt_path: str | Path,
    *,
    progress_callback: (
        Callable[[int, int, str], None]
        | None
    ) = None,
    cancel_callback: (
        Callable[[], bool]
        | None
    ) = None,
) -> dict[str, Any]:
    """
    从PPT中提取所有可识别的图片对象。

    返回结果中的 extracted 顺序即导入FrameDeck的顺序：
        页码 → 页面视觉位置（从上到下、从左到右）
    """
    source_path = Path(
        ppt_path
    )

    if (
        not source_path.is_file()
        or source_path.suffix.lower()
        not in SUPPORTED_PPT_EXTENSIONS
    ):
        raise ValueError(
            "请选择有效的 .pptx 或 .pptm 文件。"
        )

    presentation = Presentation(
        str(source_path)
    )
    slide_count = len(
        presentation.slides
    )

    records, warnings = (
        _collect_picture_records(
            presentation,
            progress_callback=progress_callback,
            cancel_callback=cancel_callback,
        )
    )

    if not records:
        return {
            "format": (
                "FrameDeck PPT Image Import"
            ),
            "version": "UI-05-43A",
            "source_file": str(
                source_path
            ),
            "output_directory": "",
            "slide_count": slide_count,
            "picture_object_count": 0,
            "extracted_count": 0,
            "skipped_count": 0,
            "slides": build_slide_summaries(
                [],
                [],
                slide_count,
            ),
            "extracted": [],
            "warnings": warnings,
            "cancelled": False,
        }

    output_dir = (
        _import_directory_for(
            source_path
        )
    )
    extracted: list[
        dict[str, Any]
    ] = []
    skipped_count = 0
    total = len(records)

    try:
        for current, record in enumerate(
            records,
            start=1,
        ):
            if (
                cancel_callback
                and cancel_callback()
            ):
                raise PPTImportCancelled(
                    "用户取消PPT图片导入。"
                )

            if progress_callback:
                progress_callback(
                    current - 1,
                    max(1, total),
                    (
                        "正在提取PPT图片："
                        f"{current}/{total}"
                    ),
                )

            shape_label = _safe_name(
                record.shape_name,
                fallback="Picture",
                max_length=32,
            )
            base_name = (
                f"P{record.slide_index:03d}_"
                f"{record.visual_order:03d}_"
                f"{shape_label}"
            )
            base_path = (
                output_dir
                / base_name
            )

            output_path, image_format, warning = (
                _write_picture_blob(
                    record.blob,
                    base_path,
                )
            )

            if output_path is None:
                skipped_count += 1
                warnings.append(
                    (
                        f"第{record.slide_index}页 "
                        f"{shape_label}：{warning}"
                    )
                )
                continue

            if warning:
                warnings.append(
                    (
                        f"第{record.slide_index}页 "
                        f"{shape_label}：{warning}"
                    )
                )

            extracted.append(
                {
                    "path": str(
                        output_path
                    ),
                    "slide_index": (
                        record.slide_index
                    ),
                    "visual_order": (
                        record.visual_order
                    ),
                    "shape_name": (
                        record.shape_name
                    ),
                    "source_format": (
                        image_format
                    ),
                    "relationship_id": (
                        record.relationship_id
                    ),
                    "transform": (
                        _transform_from_record(
                            record
                        )
                    ),
                }
            )

        slide_summaries = (
            build_slide_summaries(
                records,
                extracted,
                slide_count,
            )
        )

        manifest = {
            "format": (
                "FrameDeck PPT Image Import"
            ),
            "version": "UI-05-43A",
            "source_file": str(
                source_path
            ),
            "slide_count": slide_count,
            "picture_object_count": len(
                records
            ),
            "extracted_count": len(
                extracted
            ),
            "skipped_count": skipped_count,
            "slides": slide_summaries,
            "extracted": extracted,
            "warnings": warnings,
        }

        (
            output_dir
            / "FrameDeck_PPT_Import_Manifest.json"
        ).write_text(
            json.dumps(
                manifest,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        if progress_callback:
            progress_callback(
                total,
                max(1, total),
                (
                    "PPT图片提取完成："
                    f"{len(extracted)} 张"
                ),
            )

        return {
            **manifest,
            "output_directory": str(
                output_dir
            ),
            "cancelled": False,
        }

    except PPTImportCancelled:
        shutil.rmtree(
            output_dir,
            ignore_errors=True,
        )
        return {
            "source_file": str(
                source_path
            ),
            "output_directory": "",
            "slide_count": slide_count,
            "picture_object_count": len(
                records
            ),
            "extracted_count": 0,
            "skipped_count": 0,
            "slides": build_slide_summaries(
                records,
                [],
                slide_count,
            ),
            "extracted": [],
            "warnings": warnings,
            "cancelled": True,
        }

    except Exception:
        shutil.rmtree(
            output_dir,
            ignore_errors=True,
        )
        raise


def ppt_import_diagnostic(
    ppt_path: str | Path,
) -> str:
    """
    便于命令行测试的诊断入口。
    """
    try:
        result = extract_ppt_images(
            ppt_path
        )
        return json.dumps(
            result,
            ensure_ascii=False,
            indent=2,
        )
    except Exception:
        return traceback.format_exc()
