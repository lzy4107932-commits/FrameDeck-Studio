from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PySide6.QtWidgets import QApplication

import gui.main_window as main_window_module
from core.group_layout import PAGINATION_CONTINUOUS, PAGINATION_GROUPED
from core.ppt_generator import generate_ppt
from gui.main_window import MainWindow


def _new_window(app: QApplication, paths: list[str]) -> MainWindow:
    window = MainWindow(app)
    window.autosave_timer.stop()
    window.preview_timer.stop()
    window._config_save_timer.stop()
    window._page_thumbnail_debounce_timer.stop()
    window._page_thumbnail_render_timer.stop()
    window.row_step.setValue(1)
    window.col_step.setValue(4)
    for path in paths:
        window.image_list.addItem(window._create_external_image_item(path))
        window.image_transforms[path] = window.get_transform(path)
    window._image_groups = [
        {"id": "group_a", "name": "A", "start": 0},
        {"id": "group_b", "name": "B", "start": 4},
    ]
    window._manual_page_breaks = {4}
    window._group_page_breaks = {4}
    window._set_pagination_mode(PAGINATION_GROUPED, refresh=False)
    window.refresh_preview()
    app.processEvents()
    return window


def _close_window(window: MainWindow, app: QApplication) -> None:
    window._close_in_progress = True
    window.autosave_timer.stop()
    window.preview_timer.stop()
    window._config_save_timer.stop()
    window._page_thumbnail_debounce_timer.stop()
    window._page_thumbnail_render_timer.stop()
    window.image_list.clear()
    window.deleteLater()
    app.processEvents()


def main() -> None:
    app = QApplication.instance() or QApplication([])
    original_config_file = main_window_module.CONFIG_FILE

    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        main_window_module.CONFIG_FILE = root / "config.json"
        paths = []
        for index in range(10):
            path = root / f"image_{index + 1:02d}.png"
            Image.new(
                "RGB",
                (160, 100),
                (20 + index * 20, 80, 180 - index * 10),
            ).save(path)
            paths.append(str(path))

        window = _new_window(app, paths)
        try:
            assert window._computed_page_ranges() == [(0, 4), (4, 8), (8, 10)]

            # Cancel path: changes made during Auto Layout return to the exact
            # pre-layout version, including rows and deleted image content.
            window.row_step.setValue(2)
            window.auto_layout_check.setChecked(True)
            app.processEvents()
            assert window._current_pagination_mode() == PAGINATION_CONTINUOUS
            assert window._computed_page_ranges() == [(0, 8), (8, 10)]
            window.image_list.clearSelection()
            window.image_list.item(2).setSelected(True)
            window.image_list.setCurrentRow(2)
            window.delete_selected_image()
            assert len(window.ordered_images()) == 9
            window.row_step.setValue(3)
            window.auto_layout_check.setChecked(False)
            app.processEvents()
            assert window.row_step.value() == 2
            assert window.ordered_images() == paths
            assert window._current_pagination_mode() == PAGINATION_GROUPED
            assert window._computed_page_ranges() == [(0, 4), (4, 10)]

            # Pending Auto Layout survives save/reopen, and turning it off
            # after reopening still restores the saved pre-layout version.
            window.auto_layout_check.setChecked(True)
            app.processEvents()
            window.image_list.clearSelection()
            window.image_list.item(1).setSelected(True)
            window.image_list.setCurrentRow(1)
            window.delete_selected_image()
            pending_path = root / "auto_layout_pending.fds"
            pending_path.write_text(
                json.dumps(window.project_data(), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            pending = _new_window(app, [])
            try:
                pending._apply_project_data(
                    json.loads(pending_path.read_text(encoding="utf-8")),
                    str(pending_path),
                )
                app.processEvents()
                assert pending._current_pagination_mode() == PAGINATION_CONTINUOUS
                assert pending._auto_layout_restore_state is not None
                assert len(pending.ordered_images()) == 9
                pending.auto_layout_check.setChecked(False)
                app.processEvents()
                assert pending.ordered_images() == paths
                assert pending._current_pagination_mode() == PAGINATION_GROUPED
                assert pending._computed_page_ranges() == [(0, 4), (4, 10)]
            finally:
                _close_window(pending, app)

            # Return the original window to its pre-layout version, then test
            # the separate confirm path.
            window.auto_layout_check.setChecked(False)
            app.processEvents()
            assert window.ordered_images() == paths

            # Confirm path: the continuous result becomes fixed page boundaries.
            window.auto_layout_check.setChecked(True)
            app.processEvents()
            window.confirm_auto_layout()
            app.processEvents()
            assert window._confirmed_auto_layout
            assert window._computed_page_ranges() == [(0, 8), (8, 10)]

            window.image_list.clearSelection()
            window.image_list.item(2).setSelected(True)
            window.image_list.setCurrentRow(2)
            window.delete_selected_image()
            app.processEvents()
            assert window._computed_page_ranges() == [(0, 7), (7, 9)]

            # Project JSON round trip keeps the confirmed fixed layout.
            project_path = root / "auto_layout_confirmed.fds"
            project_path.write_text(
                json.dumps(window.project_data(), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            reloaded = _new_window(app, [])
            try:
                reloaded._apply_project_data(
                    json.loads(project_path.read_text(encoding="utf-8")),
                    str(project_path),
                )
                app.processEvents()
                assert reloaded._confirmed_auto_layout
                assert reloaded._current_pagination_mode() == PAGINATION_GROUPED
                assert reloaded._computed_page_ranges() == [(0, 7), (7, 9)]

                # PPT export consumes the same effective boundaries as preview.
                output = root / "auto_layout_confirmed.pptx"
                settings = reloaded.generation_settings()
                settings["output_file"] = str(output)
                generate_ppt(**settings)
                presentation = Presentation(str(output))
                assert len(presentation.slides) == 2
                picture_counts = [
                    sum(
                        1
                        for shape in slide.shapes
                        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                    )
                    for slide in presentation.slides
                ]
                assert picture_counts == [7, 2], picture_counts
            finally:
                _close_window(reloaded, app)

            print(
                "auto layout transaction: cancel restores old version; "
                "confirm fixes page gaps; project reload and PPT export match preview"
            )
        finally:
            _close_window(window, app)
            main_window_module.CONFIG_FILE = original_config_file


if __name__ == "__main__":
    main()
