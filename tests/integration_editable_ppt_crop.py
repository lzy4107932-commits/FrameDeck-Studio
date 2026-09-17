from __future__ import annotations

import io
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.ppt_generator import generate_ppt


def main() -> None:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        source = root / "source.png"
        output = root / "editable_crop.pptx"
        image = Image.new("RGB", (800, 500), "#CC3344")
        image.paste("#3377CC", (400, 0, 800, 500))
        image.save(source)

        generate_ppt(
            [str(source)],
            str(output),
            image_transforms={
                str(source): {
                    "zoom": 1.25,
                    "offset_x": 0.08,
                    "offset_y": -0.05,
                    "rotation": 0,
                    "flip_h": False,
                    "flip_v": False,
                    "crop": {
                        "enabled": True,
                        "x": 0.20,
                        "y": 0.10,
                        "width": 0.60,
                        "height": 0.75,
                    },
                }
            },
            rows=1,
            cols=1,
            image_mode="裁剪填充",
            show_page_number=False,
        )

        presentation = Presentation(str(output))
        picture = next(
            shape
            for shape in presentation.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        assert any(
            value > 0.001
            for value in (
                picture.crop_left,
                picture.crop_top,
                picture.crop_right,
                picture.crop_bottom,
            )
        )

        with zipfile.ZipFile(output) as archive:
            media_names = [
                name for name in archive.namelist() if name.startswith("ppt/media/")
            ]
            assert len(media_names) == 1
            with Image.open(io.BytesIO(archive.read(media_names[0]))) as embedded:
                assert embedded.size == (800, 500)

        picture.crop_left = 0
        picture.crop_top = 0
        picture.crop_right = 0
        picture.crop_bottom = 0
        reset_output = root / "reset_crop.pptx"
        presentation.save(reset_output)
        reset_picture = next(
            shape
            for shape in Presentation(str(reset_output)).slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        assert reset_picture.crop_left == 0
        assert reset_picture.crop_top == 0
        assert reset_picture.crop_right == 0
        assert reset_picture.crop_bottom == 0

        rotated_output = root / "editable_rotated_crop.pptx"
        generate_ppt(
            [str(source)],
            str(rotated_output),
            image_transforms={
                str(source): {
                    "zoom": 1.0,
                    "offset_x": 0.0,
                    "offset_y": 0.0,
                    "rotation": 90,
                    "flip_h": True,
                    "flip_v": False,
                    "crop": {
                        "enabled": True,
                        "x": 0.20,
                        "y": 0.10,
                        "width": 0.60,
                        "height": 0.75,
                    },
                }
            },
            rows=1,
            cols=1,
            image_mode="裁剪填充",
            show_page_number=False,
        )
        with zipfile.ZipFile(rotated_output) as archive:
            media_name = next(
                name for name in archive.namelist() if name.startswith("ppt/media/")
            )
            with Image.open(io.BytesIO(archive.read(media_name))) as embedded:
                assert embedded.size == (500, 800)

        print(
            "editable PPT crop: full 800x500 source retained; "
            "native crop can be reset; rotated/flip conversion retains all pixels"
        )


if __name__ == "__main__":
    main()
