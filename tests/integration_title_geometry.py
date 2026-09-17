from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.ppt_generator import generate_ppt


def main() -> None:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        image_path = root / "image.png"
        output_path = root / "title_geometry.pptx"
        Image.new("RGB", (640, 480), "#4477AA").save(image_path)

        generate_ppt(
            [str(image_path)],
            str(output_path),
            show_title=True,
            title_text="Large batch title",
            title_height=0.35,
            rows=1,
            cols=1,
            show_page_number=False,
            resolved_title_pages=[
                {
                    "title": "Large batch title",
                    "style": {
                        "font_family": "Microsoft YaHei",
                        "font_size": 48,
                        "bold": True,
                        "alignment": "center",
                        "top_spacing_cm": 0,
                        "region_height_cm": 0.35,
                    },
                }
            ],
        )

        presentation = Presentation(str(output_path))
        slide = presentation.slides[0]
        title_shape = next(
            shape
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and "Large batch title" in shape.text
        )
        picture_shape = next(
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )

        assert title_shape.height > 0
        assert title_shape.top + title_shape.height <= picture_shape.top
        assert title_shape.height > 0.35 * 360000
        print(
            "title geometry: 48 pt title expanded safely; "
            "exported title box does not overlap the image"
        )


if __name__ == "__main__":
    main()
